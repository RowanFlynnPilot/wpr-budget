"""
extract_dor.py - Marathon County taxing districts from the Wisconsin Department
of Revenue's "Town, Village and City Taxes" bulletin (TVC workbook).

Every municipality's property-tax levies by government (county, municipality,
school districts, technical college, other/TIF), its full and TIF-out values,
assessment level, and the school-levy tax credit — plus, per municipality, each
school district's share of value and levy. It is the one statewide source that
covers every municipality on the same basis, so it is what lets the tool:
(Every figure is a whole dollar rounded part by part — levies from cents,
school-district values from apportioned shares — so sums of parts reconcile
within $1 per part.)

- give a City of Wausau homeowner in the D.C. EVEREST school district the right
  bill (the City book's rate table assumes the Wausau School District; ~5% of the
  city's value is in D.C. Everest, whose school rate is lower), and
- (next) give every Marathon County municipality its own tax bill.

    python scripts/extract_dor.py sources/dor-tvc25.xlsx public/marathon-tax-districts.json

Source: https://www.revenue.wi.gov/Pages/Report/tvc-bulletin.aspx (posted each
June; tvc25 = taxes levied in 2025, payable in 2026). revenue.wi.gov does not
block scripted downloads. Reconciled per municipality, within the whole-dollar
rounding above: the five levies sum to the printed total tax, and the school
districts' levies and values sum to the school tax and the TIF-out full value.
"""
import re
import argparse

from lib import write_json

COUNTY = "MARATHON COUNTY"
TYPES = {"T": "Town", "V": "Village", "C": "City"}

# Municipal Data columns, keyed by the header text stacked over each column
# (title rows 2-4 of the sheet). Checked on every run so a re-laid-out
# workbook fails loud instead of shifting values into the wrong fields.
MUNI_COLS = {
    "code": (2, "DIST CODE"),
    "name": (3, "TYPE/NAME"),
    "population": (4, "POPULATION"),
    "full_value": (5, "FULL VALUE"),
    "assessment_level": (6, "LEVEL"),
    "full_value_ex_tif": (7, "EXCLUDING TIF"),
    "total_tax": (10, "PROPERTY TAX"),
    "school_levy_credit": (11, "TAX CREDIT"),
    "gross_rate": (12, "GROSS RATE"),
    "effective_rate": (13, "EFFECTIVE RATE"),
    "school": (14, "SECONDARY SCHOOL TAX"),
    "technical_college": (15, "COLLEGE SCHOOL TAX"),
    "county": (16, "COUNTY TAX"),
    "local": (17, "LOCAL TAX"),
    "other": (18, "OTHER TAXES"),
    "territory": (19, "TERRITORY IN TAX DISTRICT"),
}
# Technical college districts, by the code the bulletin lists among the
# districts "containing territory" (the codes that aren't school districts).
# Three Marathon County towns (Day, Green Valley, McMillan) straddle both.
TECH_COLLEGES = {"1500": "Northcentral Technical College", "1400": "Mid-State Technical College"}
SD_COLS = {"sd_code": (0, "CODE"), "sd_name": (1, "SCHOOL NAME"), "code": (2, "DISTRICT CODE"),
           "full_value": (4, "WITHIN SCHOOL DISTRICT"), "levy": (6, "DISTRICT LEVY")}


def display_name(raw):
    """DOR's all-caps names -> how the places style themselves: 'MCMILLAN' ->
    'McMillan', 'D C EVEREST AREA (ROTHSCHILD)' -> 'D.C. Everest Area' (the
    parenthetical is DOR's note of the district office's town, not its name)."""
    name = re.sub(r"\s*\(.*\)\s*$", "", " ".join(raw.split())).title()
    name = re.sub(r"\bMc([a-z])", lambda m: "Mc" + m.group(1).upper(), name)
    return re.sub(r"\bD C\b", "D.C.", name)


def headers(rows, first, last):
    """Each column's stacked header text from title rows [first, last)."""
    width = max(len(r) for r in rows[first:last])
    return [" ".join(" ".join(str(r[c] or "").split()) for r in rows[first:last] if c < len(r)).upper()
            for c in range(width)]


def check_headers(head, spec, sheet):
    for field, (col, needle) in spec.items():
        if needle not in head[col]:
            raise ValueError(f"{sheet}: column {col} ('{head[col]}') no longer carries '{needle}' ({field}) — "
                             f"the workbook layout changed; re-map MUNI_COLS/SD_COLS")


def extract(xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    mrows = list(wb["Municipal Data"].iter_rows(values_only=True))
    title = str(mrows[0][0])
    my = re.match(r"\s*(20\d{2}) Town, Village and City Taxes", title)
    if not my:
        raise ValueError(f"unexpected bulletin title {title!r}")
    tax_year = int(my.group(1))
    check_headers(headers(mrows, 1, 4), MUNI_COLS, "Municipal Data")

    munis = {}
    for r in mrows[5:]:
        if not r[0] or str(r[0]).strip() != COUNTY:
            continue
        code = str(r[2]).strip().strip('"')
        kind, _, name = str(r[3]).strip().partition(" ")
        if kind not in TYPES:
            raise ValueError(f"municipality {r[3]!r}: unknown type prefix {kind!r}")
        taxes = {k: int(r[MUNI_COLS[k][0]]) for k in ("county", "local", "school", "technical_college", "other")}
        total = int(r[MUNI_COLS["total_tax"][0]])
        # The bulletin rounds each cent-level levy (the Statement of Taxes carries
        # cents) to whole dollars on its own, so the parts may miss the rounded
        # total by up to $1 per part (Town of Bergen 2025: $1).
        if abs(sum(taxes.values()) - total) > len(taxes):
            raise ValueError(f"{r[3].strip()}: levies sum to {sum(taxes.values()):,} vs total tax {total:,}")
        munis[code] = {
            "code": code,
            "type": TYPES[kind],
            "name": display_name(name),
            "split_county": bool(str(r[1]).strip()),
            "population": int(r[MUNI_COLS["population"][0]]),
            "full_value": int(r[MUNI_COLS["full_value"][0]]),
            "full_value_ex_tif": int(r[MUNI_COLS["full_value_ex_tif"][0]]),
            "assessment_level": float(r[MUNI_COLS["assessment_level"][0]]),
            "total_tax": total,
            "school_levy_credit": int(r[MUNI_COLS["school_levy_credit"][0]]),
            "gross_rate": float(r[MUNI_COLS["gross_rate"][0]]),
            "effective_rate": float(r[MUNI_COLS["effective_rate"][0]]),
            "taxes": taxes,
            "school_districts": [],
            "technical_colleges": [],
            "_territory": str(r[MUNI_COLS["territory"][0]]).split(),
        }
    if not munis:
        raise ValueError(f"no {COUNTY} rows in the Municipal Data sheet")

    srows = list(wb["School District Data"].iter_rows(values_only=True))
    check_headers(headers(srows, 1, 4), SD_COLS, "School District Data")
    for r in srows[5:]:
        if r[2] is None or str(r[2]).strip() not in munis:
            continue
        munis[str(r[2]).strip()]["school_districts"].append({
            "code": str(r[0]).strip().zfill(4),
            "name": display_name(re.sub(r"^SCH D OF\s+", "", str(r[1]).strip())),
            "full_value": int(r[4]),
            "levy": int(r[6]),
        })

    for m in munis.values():
        sds = m["school_districts"]
        label = f"{m['type']} of {m['name']}"
        if not sds:
            raise ValueError(f"{label}: no school districts in the School District Data sheet")
        # Same whole-dollar rounding as the levies above: up to $1 per district.
        if abs(sum(s["levy"] for s in sds) - m["taxes"]["school"]) > len(sds):
            raise ValueError(f"{label}: school-district levies sum to {sum(s['levy'] for s in sds):,} "
                             f"vs school tax {m['taxes']['school']:,}")
        # District values are apportioned shares of the TIF-out value, each rounded
        # to the dollar (Town of Day 2025: $1 over) — the same $1-per-part bound.
        if abs(sum(s["full_value"] for s in sds) - m["full_value_ex_tif"]) > len(sds):
            raise ValueError(f"{label}: school-district values sum to {sum(s['full_value'] for s in sds):,} "
                             f"vs TIF-out full value {m['full_value_ex_tif']:,}")
        sds.sort(key=lambda s: -s["full_value"])
        # Every listed district must be one of this municipality's school
        # districts or a known technical college; at least one college.
        territory = m.pop("_territory")
        colleges = [c for c in territory if c not in {s["code"] for s in sds}]
        unknown = [c for c in colleges if c not in TECH_COLLEGES]
        if unknown or not colleges:
            raise ValueError(f"{label}: district codes {territory} leave technical colleges {colleges} "
                             f"(unknown: {unknown}) — add to TECH_COLLEGES")
        m["technical_colleges"] = [{"code": c, "name": TECH_COLLEGES[c]} for c in sorted(colleges, reverse=True)]
        # DOR lists a district with no value in some towns (Town of Hamburg 2025:
        # Merrill Area, $0 value, $0 levy) — no home is in it, and a 0/0 rate
        # would render as $NaN, so it isn't offered. (Dropped only after the
        # colleges are derived above, which subtracts every listed school code.)
        m["school_districts"] = sds = [s for s in sds if s["full_value"] > 0]

        # Method check: the tool's bill rate (billRates in src/ui.jsx — levies
        # over the TIF-out value, the school district's levy over its value, the
        # school levy credit over ALL value) must reproduce DOR's own published
        # effective rate wherever the municipality has a single school district.
        # (DOR prints that rate to 5 decimals and its parts are whole dollars,
        # hence one cent per $1,000; the Village of Unity's sliver in Marathon
        # County is the loosest, at $0.006.)
        if len(sds) == 1:
            t = m["taxes"]
            ours = ((t["county"] + t["local"] + t["technical_college"]) / m["full_value_ex_tif"]
                    + sds[0]["levy"] / sds[0]["full_value"] - m["school_levy_credit"] / m["full_value"]) * 1000
            if abs(ours - m["effective_rate"] * 1000) > 0.01:
                raise ValueError(f"{label}: bill rate ${ours:.3f}/$1,000 vs DOR effective rate "
                                 f"${m['effective_rate'] * 1000:.2f} — the rate method no longer matches DOR's")

    return {
        "meta": {
            "source": f"Wisconsin Department of Revenue, {tax_year} Town, Village and City Taxes",
            "source_url": "https://www.revenue.wi.gov/Pages/Report/tvc-bulletin.aspx",
            "tax_year": tax_year,
            "county": "Marathon",
        },
        "municipalities": sorted(munis.values(), key=lambda m: (m["type"] != "City", m["type"] != "Village", m["name"])),
    }


def main():
    ap = argparse.ArgumentParser(description="Extract Marathon County taxing districts from the DOR TVC workbook.")
    ap.add_argument("xlsx", help="the DOR Town, Village and City Taxes workbook (tvcYY.xlsx)")
    ap.add_argument("out", help="output JSON path (public/marathon-tax-districts.json)")
    args = ap.parse_args()
    data = extract(args.xlsx)
    write_json(args.out, data)
    ms = data["municipalities"]
    multi = [f"{m['type']} of {m['name']}" for m in ms if len(m["school_districts"]) > 1]
    print(f"wrote {args.out}: {len(ms)} Marathon County municipalities, tax year {data['meta']['tax_year']} | "
          f"{len(multi)} span more than one school district")


if __name__ == "__main__":
    main()
